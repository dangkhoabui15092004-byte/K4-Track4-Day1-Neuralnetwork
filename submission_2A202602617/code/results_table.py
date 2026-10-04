"""results_table.py — Hoàn thiện lưu kết quả và điền bảng Excel.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl


def _to_serializable(val):
    """Chuyển tensor/numpy thành kiểu chuẩn Python để lưu JSON."""
    if hasattr(val, "item"):
        return val.item()
    if isinstance(val, (list, tuple)):
        return [_to_serializable(x) for x in val]
    if isinstance(val, dict):
        return {k: _to_serializable(v) for k, v in val.items()}
    return val


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    exp_id = result.get("cfg", {}).get("exp_id", "experiment")
    out_file = Path(results_dir) / f"{exp_id}.json"

    data_to_save = {
        "cfg": _to_serializable(result.get("cfg", {})),
        "history": _to_serializable(result.get("history", {})),
        "summary": _to_serializable(result.get("summary", {})),
    }

    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    return str(out_file)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    p = Path(results_dir)
    if not p.exists():
        return []

    results = []
    for f in sorted(p.glob("*.json")):
        with open(f, "r", encoding="utf-8") as fp:
            data = json.load(fp)
            results.append(data)

    results.sort(key=lambda r: r.get("cfg", {}).get("exp_id", ""))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result.get("cfg", {})
    summary = result.get("summary", {})
    exp_id = cfg.get("exp_id", "")

    # Format hidden architecture representation
    hidden_val = cfg.get("hidden", "")
    if isinstance(hidden_val, (list, tuple)):
        hidden_str = "-".join(map(str, hidden_val))
    else:
        hidden_str = str(hidden_val)

    # Format clip_norm
    clip_norm_val = cfg.get("clip_norm", None)
    clip_norm_str = "none" if clip_norm_val is None else clip_norm_val

    # Eval metrics
    eval_acc = None
    eval_macro_f1 = None
    if eval_scores is not None:
        eval_acc = eval_scores.get("accuracy", eval_scores.get("eval_acc", None))
        eval_macro_f1 = eval_scores.get("macro_f1", eval_scores.get("eval_macro_f1", None))

    # Notes
    combined_notes = notes or cfg.get("notes", "")

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce"),
        "optimizer": cfg.get("optimizer", "sgd_momentum"),
        "lr": cfg.get("lr", ""),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": hidden_str,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": clip_norm_str,
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss", None),
        "best_val_loss": summary.get("best_val_loss", None),
        "best_epoch": summary.get("best_epoch", None),
        "final_train_loss": summary.get("final_train_loss", None),
        "final_val_loss": summary.get("final_val_loss", None),
        "val_acc": summary.get("val_acc", None),
        "val_macro_f1": summary.get("val_macro_f1", None),
        "time_per_epoch_s": summary.get("time_per_epoch_s", None),
        "peak_mem_MB": summary.get("peak_mem_MB", None),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_acc if eval_acc is not None else "",
        "eval_macro_f1": eval_macro_f1 if eval_macro_f1 is not None else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": combined_notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    formula_cols = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}

    for idx, row in enumerate(rows):
        row_num = 2 + idx
        for col_idx, header in enumerate(headers, start=1):
            if header in formula_cols:
                # Nếu hàng > 61 thì tạo công thức mới, ngược lại giữ công thức có sẵn
                if row_num > 61:
                    if header == "step0_gap_vs_lnC":
                        ws.cell(row_num, col_idx, value=f'=IF(P{row_num}="","",P{row_num}-LN(7))')
                    elif header == "gap_val_minus_train":
                        ws.cell(row_num, col_idx, value=f'=IF(OR(T{row_num}="",S{row_num}=""),"",T{row_num}-S{row_num})')
                    elif header == "delta_val_f1_vs_base":
                        ws.cell(row_num, col_idx, value=f'=IF(OR(V{row_num}="",Seeds!$C$8=""),"",V{row_num}-Seeds!$C$8)')
                    elif header == "beyond_noise":
                        ws.cell(row_num, col_idx, value=f'=IF(OR(AF{row_num}="",Seeds!$C$10=""),"",IF(ABS(AF{row_num})>Seeds!$C$10,"Có","Không"))')
                continue

            if header in row:
                val = row[header]
                # Chuyển boolean hoặc None thành dạng thích hợp
                if val is None:
                    val = ""
                ws.cell(row=row_num, column=col_idx, value=val)

    wb.save(out_path)
    print(f"[results_table] Đã lưu {len(rows)} dòng vào {out_path}")
